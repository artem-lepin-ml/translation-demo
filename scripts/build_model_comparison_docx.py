"""Build the model-comparison v4 Word report from the RU + EN experiment drafts.

Concatenates report-ru-v3.md (Russian experiment report) and
paper-section-en-final.md (English paper section) into a single .docx,
converting headings, inline bold/italic/code, tables and lists, and inserting
the G6 pipeline figure where the RU report references it.

This is a plain markdown -> docx converter, not a general one: it covers only
the constructs actually present in the two source drafts (see module-level
regexes). Re-run after the sources are updated (e.g. once the qwen3.7-plus
numbers land) to refresh the .docx with no code changes:

    uv run --with python-docx python scripts/build_model_comparison_docx.py \\
        --out /path/to/model-comparison-v4.docx --final

Drop --final while qwen3.7-plus is still running to keep the "(draft)" title
marker.
"""
from __future__ import annotations

import argparse
import re
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from docx.table import Table
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[1]
DRAFTS_DIR = ROOT / "docs/experiments/2026-07-05-model-comparison/drafts"
RU_SOURCE = DRAFTS_DIR / "report-ru-v3.md"
EN_SOURCE = DRAFTS_DIR / "paper-section-en-final.md"
PIPELINE_FIGURE = DRAFTS_DIR / "g6-pipeline.png"

TITLE_BASE = "Model comparison: NER + Wikidata grounding — v4"

# --- markdown block grammar (covers exactly what the two drafts use) -------

HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
TABLE_ROW_RE = re.compile(r"^\|(.*)\|\s*$")
BULLET_RE = re.compile(r"^-\s+(.*)$")
NUMBER_RE = re.compile(r"^(\d+)\.\s+(.*)$")
FOOTNOTE_DEF_RE = re.compile(r"^\[\^([\w-]+)\]:\s*(.*)$")
SEPARATOR_CELL_RE = re.compile(r"^:?-{1,}:?$")

# Inline tokens, in priority order via alternation (bold before italic so a
# "**...**" span is never mistaken for two "*...*" spans).
INLINE_TOKEN_RE = re.compile(r"(\*\*[^*]+?\*\*|\*[^*]+?\*|`[^`]+?`|\[\^[\w-]+\])")
FOOTNOTE_REF_RE = re.compile(r"^\[\^([\w-]+)\]$")

# Where the RU report calls out the pipeline figure ("рис. 1"); the EN
# section's parenthetical "(Figure 1)" mention is intentionally not a second
# trigger — the image is inserted once, at its first reference.
IMAGE_TRIGGER_RE = re.compile(r"рис\.\s*1|g6-pipeline", re.IGNORECASE)


@dataclass
class Block:
    kind: str  # heading | para | list_bullet | list_number | table | footnote_def
    level: int = 0
    marker: str = ""
    text: str = ""
    rows: list[list[str]] = field(default_factory=list)


def _starts_new_block(line: str) -> bool:
    s = line.rstrip()
    return bool(
        HEADING_RE.match(s)
        or TABLE_ROW_RE.match(s)
        or BULLET_RE.match(s)
        or NUMBER_RE.match(s)
        or FOOTNOTE_DEF_RE.match(s)
    )


def _split_table_row(line: str) -> list[str]:
    inner = line.strip()
    if inner.startswith("|"):
        inner = inner[1:]
    if inner.endswith("|"):
        inner = inner[:-1]
    return [cell.strip() for cell in inner.split("|")]


def parse_blocks(text: str) -> list[Block]:
    """Group raw markdown lines into blocks, reflowing hard-wrapped paragraphs
    and list items (the EN draft wraps prose at ~80 columns; the RU draft
    does not) so inline bold/italic spans that straddle a line break parse
    correctly once joined."""
    lines = text.splitlines()
    blocks: list[Block] = []
    i, n = 0, len(lines)
    while i < n:
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue

        m = HEADING_RE.match(line)
        if m:
            blocks.append(Block(kind="heading", level=len(m.group(1)), text=m.group(2).strip()))
            i += 1
            continue

        if TABLE_ROW_RE.match(line):
            rows = []
            while i < n and TABLE_ROW_RE.match(lines[i].rstrip()):
                rows.append(_split_table_row(lines[i].rstrip()))
                i += 1
            blocks.append(Block(kind="table", rows=rows))
            continue

        m = FOOTNOTE_DEF_RE.match(line)
        if m:
            marker, buf = m.group(1), [m.group(2)]
            i += 1
            while i < n and lines[i].strip() and not _starts_new_block(lines[i]):
                buf.append(lines[i].strip())
                i += 1
            blocks.append(Block(kind="footnote_def", marker=marker, text=" ".join(buf).strip()))
            continue

        m = BULLET_RE.match(line)
        if m:
            buf = [m.group(1)]
            i += 1
            while i < n and lines[i].strip() and not _starts_new_block(lines[i]):
                buf.append(lines[i].strip())
                i += 1
            blocks.append(Block(kind="list_bullet", text=" ".join(buf).strip()))
            continue

        m = NUMBER_RE.match(line)
        if m:
            marker, buf = m.group(1), [m.group(2)]
            i += 1
            while i < n and lines[i].strip() and not _starts_new_block(lines[i]):
                buf.append(lines[i].strip())
                i += 1
            blocks.append(Block(kind="list_number", marker=marker, text=" ".join(buf).strip()))
            continue

        buf = [line.strip()]
        i += 1
        while i < n and lines[i].strip() and not _starts_new_block(lines[i]):
            buf.append(lines[i].strip())
            i += 1
        blocks.append(Block(kind="para", text=" ".join(buf).strip()))

    return blocks


# --- inline formatting ------------------------------------------------------


def add_inline_runs(
    paragraph: Paragraph, text: str, footnotes: dict[str, int], counter: list[int]
) -> None:
    """Render **bold**, *italic*, `code` and [^id] footnote refs as Word runs.
    Anything else (including \\placeholder{...} / ⟦...⟧ draft markers) is left
    as plain text verbatim — that's exactly what makes them visible for the
    verification pass."""
    for token in INLINE_TOKEN_RE.split(text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**") and len(token) > 4:
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`") and token.endswith("`") and len(token) > 2:
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
        elif FOOTNOTE_REF_RE.match(token):
            marker = FOOTNOTE_REF_RE.match(token).group(1)
            if marker not in footnotes:
                counter[0] += 1
                footnotes[marker] = counter[0]
            run = paragraph.add_run(f"[{footnotes[marker]}]")
            run.font.superscript = True
        elif token.startswith("*") and token.endswith("*") and len(token) > 2:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        else:
            paragraph.add_run(token)


# --- block rendering ---------------------------------------------------------


def _is_separator_row(cells: list[str]) -> bool:
    non_empty = [c for c in cells if c]
    return bool(non_empty) and all(SEPARATOR_CELL_RE.match(c) for c in non_empty)


def _table_header_and_body(rows: list[list[str]]) -> tuple[list[str], list[list[str]]]:
    header = rows[0]
    body = [r for r in rows[1:] if not _is_separator_row(r)]
    return header, body


def _compute_column_widths(
    header: list[str], body: list[list[str]], usable_width: int
) -> list[int]:
    ncols = len(header)
    weights = [max(len(header[i]), 3) for i in range(ncols)]
    for row in body:
        for i in range(min(ncols, len(row))):
            weights[i] = max(weights[i], len(row[i]))
    total = sum(weights)
    min_width = int(usable_width * 0.08)
    widths = [max(int(usable_width * w / total), min_width) for w in weights]
    overflow_scale = usable_width / sum(widths)
    if overflow_scale < 1.0:
        widths = [int(w * overflow_scale) for w in widths]
    return widths


def add_table_block(
    doc: Document, block: Block, footnotes: dict[str, int], counter: list[int]
) -> Table:
    header, body = _table_header_and_body(block.rows)
    ncols = len(header)
    table = doc.add_table(rows=1, cols=ncols)
    table.style = "Table Grid"
    table.autofit = False

    for i, cell in enumerate(table.rows[0].cells):
        text = header[i] if i < len(header) else ""
        add_inline_runs(cell.paragraphs[0], text, footnotes, counter)
        for run in cell.paragraphs[0].runs:
            run.font.bold = True

    for row in body:
        cells = table.add_row().cells
        for i, cell in enumerate(cells):
            text = row[i] if i < len(row) else ""
            add_inline_runs(cell.paragraphs[0], text, footnotes, counter)

    section = doc.sections[0]
    usable_width = section.page_width - section.left_margin - section.right_margin
    widths = _compute_column_widths(header, body, usable_width)
    for i, col in enumerate(table.columns):
        col.width = widths[i]
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = widths[i]

    return table


def add_heading_block(doc: Document, block: Block) -> None:
    # Headings in both drafts are plain text (verified: no *, **, ` in any
    # heading line), so a direct add_heading is sufficient here.
    doc.add_heading(block.text, level=min(max(block.level, 1), 9))


def add_para_block(
    doc: Document, block: Block, footnotes: dict[str, int], counter: list[int]
) -> Paragraph:
    p = doc.add_paragraph()
    add_inline_runs(p, block.text, footnotes, counter)
    return p


def add_bullet_block(
    doc: Document, block: Block, footnotes: dict[str, int], counter: list[int]
) -> None:
    p = doc.add_paragraph(style="List Bullet")
    add_inline_runs(p, block.text, footnotes, counter)


def add_numbered_block(
    doc: Document, block: Block, footnotes: dict[str, int], counter: list[int]
) -> None:
    # Word's built-in numbered-list style shares one running counter across
    # the whole document unless each list gets its own numbering definition
    # (abstractNum/num) — plumbing that isn't worth it here. The source
    # markdown already numbers each list correctly on its own, so we render
    # that number as literal text with a hanging indent instead of relying on
    # Word's auto-numbering, which would otherwise renumber later lists
    # starting from where the previous one left off.
    p = doc.add_paragraph(style="List Paragraph")
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.first_line_indent = Inches(-0.3)
    p.add_run(f"{block.marker}. ")
    add_inline_runs(p, block.text, footnotes, counter)


def add_footnote_def_block(
    doc: Document, block: Block, footnotes: dict[str, int], counter: list[int]
) -> None:
    if block.marker not in footnotes:
        counter[0] += 1
        footnotes[block.marker] = counter[0]
    p = doc.add_paragraph()
    marker_run = p.add_run(f"[{footnotes[block.marker]}] ")
    marker_run.italic = True
    marker_run.font.size = Pt(9)
    start = len(p.runs)
    add_inline_runs(p, block.text, footnotes, counter)
    for run in p.runs[start:]:
        run.italic = True
        run.font.size = Pt(9)


def insert_pipeline_figure(doc: Document, image_path: Path) -> None:
    section = doc.sections[0]
    usable_width = section.page_width - section.left_margin - section.right_margin
    doc.add_picture(str(image_path), width=usable_width)
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = caption.add_run("Figure 1. G6 label_first candidate-ladder pipeline.")
    run.italic = True
    run.font.size = Pt(9)


def render_source(
    doc: Document,
    blocks: list[Block],
    footnotes: dict[str, int],
    counter: list[int],
    image_path: Path,
    image_state: dict[str, bool],
) -> None:
    for block in blocks:
        if block.kind == "heading":
            add_heading_block(doc, block)
        elif block.kind == "table":
            add_table_block(doc, block, footnotes, counter)
        elif block.kind == "list_bullet":
            add_bullet_block(doc, block, footnotes, counter)
        elif block.kind == "list_number":
            add_numbered_block(doc, block, footnotes, counter)
        elif block.kind == "footnote_def":
            add_footnote_def_block(doc, block, footnotes, counter)
        else:
            add_para_block(doc, block, footnotes, counter)

        is_para = block.kind == "para"
        if not image_state["inserted"] and is_para and IMAGE_TRIGGER_RE.search(block.text):
            insert_pipeline_figure(doc, image_path)
            image_state["inserted"] = True


def build_docx(out_path: Path, draft: bool) -> None:
    doc = Document()

    title = TITLE_BASE + (" (draft)" if draft else "")
    doc.add_heading(title, level=0)
    doc.add_page_break()

    footnotes: dict[str, int] = {}
    counter = [0]
    image_state = {"inserted": False}

    ru_blocks = parse_blocks(RU_SOURCE.read_text(encoding="utf-8"))
    render_source(doc, ru_blocks, footnotes, counter, PIPELINE_FIGURE, image_state)

    if not image_state["inserted"]:
        # Defensive fallback per spec: no explicit figure reference found, so
        # drop it right after the RU report (which is where the pipeline is
        # discussed, § 2.4) rather than silently omitting it.
        insert_pipeline_figure(doc, PIPELINE_FIGURE)
        image_state["inserted"] = True

    doc.add_page_break()

    en_blocks = parse_blocks(EN_SOURCE.read_text(encoding="utf-8"))
    render_source(doc, en_blocks, footnotes, counter, PIPELINE_FIGURE, image_state)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", type=Path, required=True, help="Output .docx path")
    parser.add_argument(
        "--final",
        action="store_true",
        help="Drop the '(draft)' title marker (use once qwen3.7-plus numbers land in the sources)",
    )
    args = parser.parse_args()
    build_docx(args.out, draft=not args.final)
    print(f"wrote {args.out}")


if __name__ == "__main__":
    main()
