"""Export a document (spec 2026-07-05-export-xlsx; md re-contracted 2026-07-17).

Two formats, same underlying row selection: ``.xlsx`` (openpyxl — parallel
text with paragraph alignment and scores) and ``.md`` (the final translation
text itself, blank-line-separated paragraphs — owner contract 2026-07-17,
replacing the original table). No schema changes; both read straight off
``paragraph``/``score``.
"""
from __future__ import annotations

import io
import re
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_BG = "1F2035"
HEADER_FG = "C0CAF5"
SCORE_GREEN = "3DDC84"
SCORE_YELLOW = "E0AF68"
SCORE_RED = "F7768E"

_SLUG_RE = re.compile(r"[^a-z0-9]+")


def slugify(title: str) -> str:
    slug = _SLUG_RE.sub("-", title.lower()).strip("-")
    return slug or "document"


def _latest_aggregate(conn, paragraph_id: int) -> float | None:
    row = conn.execute(
        "SELECT aggregate FROM score WHERE paragraph_id=? AND kind IN ('seed','live') "
        "ORDER BY created_at DESC, id DESC LIMIT 1", (paragraph_id,)).fetchone()
    return row["aggregate"] if row else None


def _rows(conn, doc_id: int) -> list[tuple[int, str, str, float | None]]:
    paras = conn.execute(
        "SELECT * FROM paragraph WHERE document_id=? ORDER BY idx", (doc_id,)).fetchall()
    return [(p["idx"] + 1, p["source"], p["target"], _latest_aggregate(conn, p["id"]))
            for p in paras]


def _score_color(score: float) -> str:
    if score >= 8:
        return SCORE_GREEN
    if score >= 6:
        return SCORE_YELLOW
    return SCORE_RED


def build_xlsx(conn, doc) -> bytes:
    rows = _rows(conn, doc["id"])
    wb = Workbook()
    ws = wb.active
    ws.title = "Translation"

    headers = ["¶", f"Source · {doc['source_lang']}",
               f"Translation · {doc['target_lang']}", "Score"]
    header_fill = PatternFill("solid", fgColor=HEADER_BG)
    header_font = Font(bold=True, color=HEADER_FG)
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=text)
        cell.fill = header_fill
        cell.font = header_font
    ws.freeze_panes = "A2"

    meta = (f'"{doc["title"]}" · {doc["source_lang"]} → {doc["target_lang"]} · '
            f'exported {datetime.now(timezone.utc).strftime("%Y-%m-%d")} · Glossa-MT')
    ws.cell(row=2, column=1, value=meta)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=4)
    meta_font = Font(color="787C99")
    ws.cell(row=2, column=1).font = meta_font

    wrap_top = Alignment(wrap_text=True, vertical="top")
    for i, (idx, source, target, score) in enumerate(rows, start=3):
        ws.cell(row=i, column=1, value=idx).alignment = wrap_top
        ws.cell(row=i, column=2, value=source).alignment = wrap_top
        ws.cell(row=i, column=3, value=target or "").alignment = wrap_top
        score_cell = ws.cell(row=i, column=4, value=round(score, 1) if score is not None else None)
        score_cell.alignment = wrap_top
        if score is not None:
            score_cell.font = Font(color=_score_color(score))

    widths = [6, 58, 58, 9]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_markdown(conn, doc) -> str:
    # Owner contract (2026-07-17, replaces the original table bonus): the .md
    # export is the FINAL translation text itself — title, then one block per
    # translated paragraph separated by blank lines. No table, no source
    # column, no cell escaping; paragraphs without a translation yet are
    # skipped rather than emitted as empty blocks.
    rows = _rows(conn, doc["id"])
    blocks = [t.strip() for _idx, _source, target, _score in rows
              for t in [target or ""] if t.strip()]
    return "\n\n".join([f"# {doc['title']}", *blocks]) + "\n"
